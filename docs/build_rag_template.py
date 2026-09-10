from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

d=Document(); s=d.sections[0]; s.top_margin=Inches(.6); s.bottom_margin=Inches(.6); s.left_margin=Inches(.7); s.right_margin=Inches(.7)
for st in ['Normal','Title','Heading 1','Heading 2']:
 d.styles[st].font.name='STSong'; d.styles[st]._element.rPr.rFonts.set(qn('w:eastAsia'),'STSong')
d.styles['Normal'].font.size=Pt(10); d.styles['Title'].font.size=Pt(22); d.styles['Title'].font.bold=True; d.styles['Heading 1'].font.size=Pt(15); d.styles['Heading 1'].font.bold=True
def fmt(t):
 p=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
 for e in ['top','left','bottom','right','insideH','insideV']:
  x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); b.append(x)
 p.append(b)
def field(label,hint='',lines=2):
 t=d.add_table(rows=2,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; fmt(t); c=t.cell(0,0); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'EAF2F8'); c._tc.get_or_add_tcPr().append(sh); c.text=label; c.paragraphs[0].runs[0].bold=True; t.cell(1,0).text=(hint+'\n' if hint else '')+'\n'*lines; t.cell(1,0).paragraphs[0].runs[0].font.color.rgb=RGBColor(130,130,130); d.add_paragraph('')
def sec(title,desc=''):
 d.add_heading(title,1)
 if desc: d.add_paragraph(desc)
p=d.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('行业专家经验与规则采集模板')
p=d.add_paragraph('用于收集可审核、可检索、可导入 RAG 的专家经验、规则与红线'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_paragraph('填写说明：一份文档建议只描述一个可独立审核的知识单元。尽量使用具体对象、条件、数值、时间和证据。管理员审核发布后，内容才进入 RAG 索引。')
sec('一 基本信息')
for x in [('标题','例如：趋势突破后的回撤处理'),('知识类型','专家经验 / 规则 / 红线'),('专家姓名与角色',''),('所属行业与领域',''),('适用对象','产品、业务、品种或客户类型'),('适用标签','逗号分隔，例如：BTC/USDT, 15m, 趋势交易'),('来源与版本','来源可填专家访谈、文档或案例；版本建议 1.0'),('经验有效期','起止日期；长期有效仍需定期复核')]: field(*x)
sec('二 一句话结论','用一句话说明最希望系统记住什么。'); field('结论')
sec('三 适用条件','描述什么情况下可以使用这条经验，尽量可观察、可判断。')
for x in ['市场或业务状态','输入数据要求','时间范围或阶段','前置条件']: field(x)
sec('四 触发信号','填写触发规则的事实信号，避免“感觉”“适当”等无法判断的词。')
for x in ['触发信号 1','触发信号 2','触发信号 3']: field(x)
sec('五 建议动作','说明动作、对象、幅度或范围、执行时机。')
for x in ['动作 1','动作 2']: field(x,'执行时机：\n建议范围：\n预期目的：',3)
sec('六 禁止动作与红线','只有违反后必须停止、暂停或转人工时，才放入红线。')
for x in ['BLOCK 禁止动作','REVIEW 必须人工复核','违反后处理（PAUSED / MANUAL_CONTROL / 其他）']: field(x)
sec('七 例外与冲突处理')
for x in ['例外情况','优先级更高的规则','信息不足时如何处理','与其他规则冲突时如何处理']: field(x)
sec('八 失效条件')
for x in ['数据过期条件','市场或业务状态变化','指标或阈值失效','复核日期或触发事件']: field(x)
sec('九 正例与反例')
field('正例','背景：\n观察到的信号：\n采取的动作：\n结果：\n为什么适用：',5); field('反例','背景：\n看似相似但不适用的信号：\n不应采取的动作：\n原因：',4)
sec('十 证据与可信度')
for x in ['证据链接或文件名','数据区间或案例编号','样本数量与覆盖范围','专家置信度（高 / 中 / 低）','已知局限','建议复核周期']: field(x)
sec('十一 给系统的简短摘要','100 字以内，用于列表展示和人工审核。'); field('摘要')
sec('十二 管理员审核记录')
for x in ['审核结论（通过 / 退回补充 / 不采用）','审核人','审核日期','修改意见']: field(x)
d.sections[0].footer.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER; d.sections[0].footer.paragraphs[0].text='Axiom RAG 知识采集模板'
d.save('docs/rag-expert-collection-template.docx')
