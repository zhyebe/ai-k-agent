from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

d=Document(); s=d.sections[0]; s.top_margin=Inches(.7); s.bottom_margin=Inches(.7); s.left_margin=Inches(.8); s.right_margin=Inches(.8)
for n in ['Normal','Title','Heading 1']:
    d.styles[n].font.name='STSong'; d.styles[n]._element.rPr.rFonts.set(qn('w:eastAsia'),'STSong')
d.styles['Normal'].font.size=Pt(11); d.styles['Title'].font.size=Pt(22); d.styles['Title'].font.bold=True; d.styles['Heading 1'].font.size=Pt(14); d.styles['Heading 1'].font.bold=True
def q(i,title,helptext):
    d.add_heading(f'{i}. {title}',1); d.add_paragraph(helptext)
    for _ in range(7): d.add_paragraph('________________________________________________________________________________')
    d.add_paragraph('')
p=d.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('行业专家经验分享问卷')
p=d.add_paragraph('请用您平时说话和写作的方式回答，不需要使用专业的系统术语。'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_paragraph('我们会在收集后，再由智能体整理为知识库内容。可以写文字，也可以先口述后整理。没有标准答案，真实、具体的经验最有价值。')
d.add_paragraph('姓名：____________________    行业/岗位：____________________    日期：____________________')
q(1,'您主要擅长什么？','请介绍您的行业、岗位，以及您最熟悉的工作内容。')
q(2,'请分享一条您认为有价值的经验','可以写判断方法、处理方式、工作诀窍，或者一个经常提醒别人的注意事项。')
q(3,'通常遇到什么情况时，您会使用这个经验？','请描述当时的背景、现象或问题。')
q(4,'您一般会怎么判断和处理？','按照您的实际习惯描述即可，可以分步骤，也可以直接讲故事。')
q(5,'有没有不能这样做的情况？','请写例外情况、风险，或者需要特别小心的地方。')
q(6,'请分享一个真实案例','成功或失败都可以。请尽量说明当时发生了什么、您做了什么、最后结果怎样。')
q(7,'这条经验是否有时间、行业或对象限制？','例如只适用于某类客户、某种业务、某个阶段，或者只在特定条件下有效。')
q(8,'还有哪些补充说明？','任何您认为有帮助的背景、资料、数据、案例或提醒，都可以写在这里。')
d.add_paragraph('感谢您的分享。我们会保留您的原始表达，并在整理后进行人工审核。')
d.save('docs/简明版行业专家经验分享问卷.docx')
